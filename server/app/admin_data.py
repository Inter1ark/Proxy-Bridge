"""Read models for the admin panel: stats, order and license listings, exports.

All dates shown to the owner are in Moscow time (UTC+3, no DST since 2014).
"""

import csv
import io
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .config import settings
from .models import Device, License, Order, StoreOrder, StoreProxy, from_db, iso_z, to_db, utcnow
from .plans import PLANS
from .store.service import _product_title

MSK = timezone(timedelta(hours=3))

FAIL_REASON_TEXT = {
    "expired_on_confirmation": "Покупатель не завершил оплату, время истекло",
    "country_forbidden": "Карта иностранного банка, оплата запрещена",
    "insufficient_funds": "Недостаточно средств",
    "card_expired": "Срок карты истёк",
    "fraud_suspected": "Банк заподозрил мошенничество",
    "general_decline": "Банк отклонил без объяснения",
    "3d_secure_failed": "Не пройдено подтверждение 3-D Secure",
    "canceled_by_merchant": "Отменён продавцом",
    "permission_revoked": "Отозвано разрешение на списание",
    "call_issuer": "Банк просит покупателя позвонить в банк",
    "payment_method_limit_exceeded": "Превышен лимит по способу оплаты",
    "payment_method_restricted": "Способ оплаты ограничен",
    "internal_timeout": "Технический сбой у платёжной системы",
    "expired": "Счёт истёк, оплата не поступила",
}
PENDING_TEXT = "Ожидает оплаты"

PAY_TYPE_TEXT = {
    "bank_card": "Банковская карта",
    "sbp": "СБП",
    "sberbank": "SberPay",
    "tinkoff_bank": "T-Pay",
    "yoo_money": "ЮMoney",
}

STATUS_TEXT = {"paid": "Оплачен", "canceled": "Не оплачен", "pending": "Ожидает оплаты"}
METHOD_TEXT = {"yookassa": "ЮKassa", "cryptobot": "CryptoBot"}
PERIODS = ("today", "7d", "30d", "all")


# ---------- small helpers ----------

def plan_title(plan_id: str | None) -> str:
    plan = PLANS.get(plan_id or "")
    return plan.title if plan else (plan_id or "")


def fail_reason_text(order: Order) -> str | None:
    if order.status == "pending":
        return PENDING_TEXT
    if not order.fail_reason:
        return None
    return FAIL_REASON_TEXT.get(order.fail_reason, order.fail_reason)


def pay_type_text(pay_type: str | None) -> str | None:
    if not pay_type:
        return None
    return PAY_TYPE_TEXT.get(pay_type, pay_type)


def is_test_email(email: str | None) -> bool:
    return bool(email) and email.strip().lower() in settings.test_emails


def to_msk(dt: datetime | None) -> datetime | None:
    dt = from_db(dt)
    return dt.astimezone(MSK) if dt else None


def fmt_msk(dt: datetime | None) -> str:
    m = to_msk(dt)
    return m.strftime("%d.%m.%Y %H:%M") if m else ""


def msk_day_start_utc(day: date) -> datetime:
    """Naive UTC datetime of 00:00 Moscow time on the given day."""
    return to_db(datetime(day.year, day.month, day.day, tzinfo=MSK))


def parse_day(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError:
        return None


def short_hwid(hwid: str | None) -> str:
    return (hwid or "")[:10]


def license_state(lic: License, now: datetime | None = None) -> str:
    if lic.revoked:
        return "revoked"
    if lic.is_expired(now):
        return "expired"
    return "active"


def device_dict(d: Device) -> dict:
    return {
        "hwid": short_hwid(d.hwid),
        "app_version": d.app_version,
        "activated_at": iso_z(d.activated_at),
        "last_seen": iso_z(d.last_seen),
    }


def license_dict(lic: License, with_devices: bool = True) -> dict:
    data = {
        "key": lic.key,
        "plan_id": lic.plan_id,
        "plan_title": plan_title(lic.plan_id),
        "email": lic.email,
        "source": lic.source,
        "order_id": lic.order_id,
        "issued_at": iso_z(lic.issued_at),
        "expires_at": iso_z(lic.expires_at),
        "revoked": bool(lic.revoked),
        "state": license_state(lic),
        "is_test": is_test_email(lic.email),
        "device_count": len(lic.devices),
    }
    if with_devices:
        devices = sorted(lic.devices, key=lambda d: d.last_seen or d.activated_at, reverse=True)
        data["devices"] = [device_dict(d) for d in devices]
    return data


# ---------- orders ----------

def order_item(order: Order, lic_info: dict | None = None) -> dict:
    lic_info = lic_info or {}
    return {
        "id": order.id,
        "plan_id": order.plan_id,
        "plan_title": plan_title(order.plan_id),
        "email": order.email,
        "method": order.method,
        "method_title": METHOD_TEXT.get(order.method, order.method),
        "status": order.status,
        "status_text": STATUS_TEXT.get(order.status, order.status),
        "amount_rub": order.amount_rub,
        "income_rub": order.income_rub,
        "provider_payment_id": order.provider_payment_id,
        "key": order.key,
        "created_at": iso_z(order.created_at),
        "paid_at": iso_z(order.paid_at),
        "checked_at": iso_z(order.checked_at),
        "fail_reason": order.fail_reason,
        "fail_reason_text": fail_reason_text(order),
        "pay_type": order.pay_type,
        "pay_type_text": pay_type_text(order.pay_type),
        "device_count": lic_info.get("device_count", 0),
        "last_seen": lic_info.get("last_seen"),
        "is_test": is_test_email(order.email),
    }


def license_info_for_orders(db: Session, orders: list[Order]) -> dict[int, dict]:
    """order id -> {"device_count", "last_seen"} for the license issued for each order."""
    by_key = {o.key: o.id for o in orders if o.key}
    order_ids = [o.id for o in orders]
    if not order_ids:
        return {}
    conds = [License.order_id.in_(order_ids)]
    if by_key:
        conds.append(License.key.in_(list(by_key)))
    rows = db.execute(
        select(License.id, License.key, License.order_id,
               func.count(Device.id), func.max(Device.last_seen))
        .select_from(License)
        .outerjoin(Device, Device.license_id == License.id)
        .where(or_(*conds))
        .group_by(License.id)
    ).all()
    result: dict[int, dict] = {}
    for _lid, key, order_id, count, last_seen in rows:
        oid = order_id if order_id in order_ids else by_key.get(key)
        if oid is not None:
            result[oid] = {"device_count": int(count or 0), "last_seen": iso_z(last_seen)}
    return result


def order_filters(status: str | None = None, method: str | None = None, plan: str | None = None,
                  q: str | None = None, date_from: str | None = None, date_to: str | None = None,
                  exclude_test: bool = False) -> list:
    conds = []
    if status:
        conds.append(Order.status == status)
    if method:
        conds.append(Order.method == method)
    if plan:
        conds.append(Order.plan_id == plan)
    q = (q or "").strip()
    if q:
        like = f"%{q.lower()}%"
        parts = [func.lower(Order.email).like(like),
                 func.lower(Order.key).like(like),
                 func.lower(Order.provider_payment_id).like(like)]
        digits = q.lstrip("#")
        if digits.isdigit():
            parts.append(Order.id == int(digits))
        conds.append(or_(*parts))
    d_from, d_to = parse_day(date_from), parse_day(date_to)
    if d_from:
        conds.append(Order.created_at >= msk_day_start_utc(d_from))
    if d_to:
        conds.append(Order.created_at < msk_day_start_utc(d_to + timedelta(days=1)))
    if exclude_test and settings.test_emails:
        conds.append(func.lower(Order.email).notin_(sorted(settings.test_emails)))
    return conds


def list_orders(db: Session, conds: list, page: int = 1, per_page: int = 50) -> dict:
    total = db.scalar(select(func.count(Order.id)).where(*conds)) or 0
    orders = db.scalars(select(Order).where(*conds).order_by(Order.id.desc())
                        .offset((page - 1) * per_page).limit(per_page)).all()
    info = license_info_for_orders(db, orders)
    return {
        "items": [order_item(o, info.get(o.id)) for o in orders],
        "total": total,
        "page": page,
        "per_page": per_page,
        "pages": max(1, -(-total // per_page)),
    }


def order_details(db: Session, order: Order) -> dict:
    lic = db.scalar(select(License).where(License.order_id == order.id))
    if lic is None and order.key:
        lic = db.scalar(select(License).where(License.key == order.key))
    info = None
    if lic is not None:
        last = max((d.last_seen for d in lic.devices if d.last_seen), default=None)
        info = {"device_count": len(lic.devices), "last_seen": iso_z(last)}
    data = order_item(order, info)
    data["license"] = license_dict(lic) if lic is not None else None
    return data


# ---------- stats ----------

def _empty_bucket() -> dict:
    return {"orders": 0, "paid": 0, "canceled": 0, "pending": 0, "conversion": 0.0,
            "gross": 0, "net": 0.0, "net_known": 0, "by_plan": {}, "by_method": {}, "by_pay_type": {}}


def _add_breakdown(target: dict, key: str, title: str, amount: int) -> None:
    row = target.setdefault(key, {"id": key, "title": title, "count": 0, "gross": 0})
    row["count"] += 1
    row["gross"] += amount


def compute_stats(db: Session, exclude_test: bool = False, now: datetime | None = None) -> dict:
    now = now or utcnow()
    now_msk = now.astimezone(MSK)
    today = now_msk.date()
    starts = {
        "today": msk_day_start_utc(today),
        "7d": msk_day_start_utc(today - timedelta(days=6)),
        "30d": msk_day_start_utc(today - timedelta(days=29)),
        "all": None,
    }
    test_emails = settings.test_emails if exclude_test else set()

    buckets = {p: _empty_bucket() for p in PERIODS}
    days = [today - timedelta(days=i) for i in range(29, -1, -1)]
    by_day = {d: {"date": d.isoformat(), "paid_count": 0, "gross": 0} for d in days}

    rows = db.execute(select(Order.status, Order.plan_id, Order.method, Order.amount_rub,
                             Order.income_rub, Order.pay_type, Order.email, Order.created_at)).all()
    for status, plan_id, method, amount, income, pay_type, email, created_at in rows:
        if test_emails and (email or "").strip().lower() in test_emails:
            continue
        amount = int(amount or 0)
        for period, start in starts.items():
            if start is not None and (created_at is None or created_at < start):
                continue
            b = buckets[period]
            b["orders"] += 1
            if status in ("paid", "canceled", "pending"):
                b[status] += 1
            if status == "paid":
                b["gross"] += amount
                if income is not None:
                    b["net"] += float(income)
                    b["net_known"] += 1
                else:
                    b["net"] += amount
                _add_breakdown(b["by_plan"], plan_id, plan_title(plan_id), amount)
                _add_breakdown(b["by_method"], method, METHOD_TEXT.get(method, method), amount)
                pt = pay_type or "unknown"
                _add_breakdown(b["by_pay_type"], pt, pay_type_text(pay_type) or "Нет данных", amount)
        if status == "paid" and created_at is not None:
            day = to_msk(created_at).date()
            if day in by_day:
                by_day[day]["paid_count"] += 1
                by_day[day]["gross"] += amount

    for b in buckets.values():
        b["conversion"] = round(b["paid"] * 100.0 / b["orders"], 1) if b["orders"] else 0.0
        b["net"] = round(b["net"], 2)
        for name in ("by_plan", "by_method", "by_pay_type"):
            b[name] = sorted(b[name].values(), key=lambda r: (-r["gross"], -r["count"]))

    # Licenses (not tied to a period).
    now_db = to_db(now)
    lic_conds = []
    if test_emails:
        lic_conds.append(or_(License.email.is_(None), func.lower(License.email).notin_(sorted(test_emails))))
    active_cond = [License.revoked.is_(False),
                   or_(License.expires_at.is_(None), License.expires_at > now_db)]
    active = db.scalar(select(func.count(License.id)).where(*lic_conds, *active_cond)) or 0
    total_lic = db.scalar(select(func.count(License.id)).where(*lic_conds)) or 0
    with_devices = db.scalar(
        select(func.count(func.distinct(Device.license_id)))
        .join(License, License.id == Device.license_id).where(*lic_conds)) or 0
    activations_24h = db.scalar(
        select(func.count(Device.id))
        .join(License, License.id == Device.license_id)
        .where(*lic_conds, Device.activated_at >= to_db(now - timedelta(hours=24)))) or 0

    return {
        "now": iso_z(now),
        "today_msk": today.isoformat(),
        "exclude_test": bool(exclude_test),
        "test_emails": sorted(settings.test_emails),
        "periods": buckets,
        "licenses": {"total": total_lic, "active": active, "with_devices": with_devices,
                     "activations_24h": activations_24h},
        "revenue_by_day": [by_day[d] for d in days],
    }


# ---------- licenses ----------

def license_filters(q: str | None = None, state: str | None = None, exclude_test: bool = False,
                    now: datetime | None = None) -> list:
    now_db = to_db(now or utcnow())
    conds = []
    q = (q or "").strip()
    if q:
        like = f"%{q.lower()}%"
        parts = [func.lower(License.key).like(like), func.lower(License.email).like(like)]
        digits = q.lstrip("#")
        if digits.isdigit():
            parts.append(License.order_id == int(digits))
        conds.append(or_(*parts))
    if state == "active":
        conds += [License.revoked.is_(False), or_(License.expires_at.is_(None), License.expires_at > now_db)]
    elif state == "revoked":
        conds.append(License.revoked.is_(True))
    elif state == "expired":
        conds += [License.revoked.is_(False), License.expires_at.isnot(None), License.expires_at <= now_db]
    if exclude_test and settings.test_emails:
        conds.append(or_(License.email.is_(None), func.lower(License.email).notin_(sorted(settings.test_emails))))
    return conds


def list_licenses(db: Session, conds: list, page: int = 1, per_page: int = 50) -> dict:
    total = db.scalar(select(func.count(License.id)).where(*conds)) or 0
    lics = db.scalars(select(License).where(*conds).order_by(License.id.desc())
                      .offset((page - 1) * per_page).limit(per_page)).all()
    return {
        "items": [license_dict(lic) for lic in lics],
        "total": total,
        "page": page,
        "per_page": per_page,
        "pages": max(1, -(-total // per_page)),
    }


# ---------- proxy store ----------

STORE_STATUSES = ("pending", "paid", "canceled")
STORE_FULFILL = ("queued", "working", "done", "failed", "manual")
STORE_PROXY_STATUSES = ("provisioning", "active", "exhausted", "expired", "failed")
ATTENTION_FULFILL = ("failed", "manual")


def store_needs_attention(order: StoreOrder) -> bool:
    return order.fulfill in ATTENTION_FULFILL and not order.refunded


# Only "failed" means "nothing was bought". A "manual" order may have reached the vendor
# (purchase outcome unknown), so it is never retried from the panel.
RETRY_FULFILL = ("failed",)


def store_can_retry(order: StoreOrder) -> bool:
    """Nothing was bought at the vendor, so the worker may simply try again."""
    return (order.status == "paid" and order.fulfill in RETRY_FULFILL and not order.refunded
            and not order.vendor_order_id)


def store_order_item(order: StoreOrder) -> dict:
    try:
        title = _product_title(order)
    except (ValueError, TypeError):
        title = order.product
    return {
        "id": order.id,
        "sid": f"S{order.id}",
        "created_at": iso_z(order.created_at),
        "paid_at": iso_z(order.paid_at),
        "license_key": order.license_key,
        "product": order.product,
        "title": title,
        "amount_rub": order.amount_rub,
        "income_rub": order.income_rub,
        "cost_usd": order.cost_usd,
        "method": order.method,
        "method_title": METHOD_TEXT.get(order.method, order.method),
        "pay_type": order.pay_type,
        "pay_type_text": pay_type_text(order.pay_type),
        "status": order.status,
        "fulfill": order.fulfill,
        "fulfill_error": order.fulfill_error,
        "refunded": bool(order.refunded),
        "proxy_id": order.proxy_id,
        "attention": store_needs_attention(order),
        "can_retry": store_can_retry(order),
    }


def store_order_filters(status: str | None = None, fulfill: str | None = None, q: str | None = None) -> list:
    """fulfill: one state, a comma separated list, or "attention" (failed or manual, not refunded)."""
    conds = []
    if status:
        conds.append(StoreOrder.status == status)
    if fulfill == "attention":
        conds += [StoreOrder.fulfill.in_(ATTENTION_FULFILL), StoreOrder.refunded.is_(False)]
    elif fulfill:
        conds.append(StoreOrder.fulfill.in_([f for f in fulfill.split(",") if f]))
    q = (q or "").strip()
    if q:
        like = f"%{q.lower()}%"
        parts = [func.lower(StoreOrder.license_key).like(like), StoreOrder.token == q]
        digits = q.lstrip("#").lstrip("Ss")
        if digits.isdigit():
            parts.append(StoreOrder.id == int(digits))
        conds.append(or_(*parts))
    return conds


def list_store_orders(db: Session, conds: list, limit: int = 200) -> dict:
    total = db.scalar(select(func.count(StoreOrder.id)).where(*conds)) or 0
    orders = db.scalars(select(StoreOrder).where(*conds).order_by(StoreOrder.id.desc()).limit(limit)).all()
    return {"items": [store_order_item(o) for o in orders], "total": total, "limit": limit}


def store_proxy_item(p: StoreProxy) -> dict:
    # Never include the proxy login or password here.
    return {
        "id": p.id,
        "license_key": p.license_key,
        "kind": p.kind,
        "vendor": p.vendor,
        "status": p.status,
        "country_code": p.country_code,
        "state": p.state,
        "city": p.city,
        "ptype": p.ptype,
        "rotation": p.rotation,
        "ttl": p.ttl,
        "gb_total": p.gb_total,
        "gb_used": round((p.bytes_used or 0) / 1e9, 3),
        "expires_at": iso_z(p.expires_at),
        "renew_pending": int(p.renew_pending or 0),
        "created_at": iso_z(p.created_at),
        "address": f"{p.host}:{p.port}" if p.host else "",
    }


def store_proxy_filters(status: str | None = None, q: str | None = None) -> list:
    conds = []
    if status:
        conds.append(StoreProxy.status == status)
    q = (q or "").strip()
    if q:
        like = f"%{q.lower()}%"
        parts = [func.lower(StoreProxy.license_key).like(like), func.lower(StoreProxy.host).like(like)]
        digits = q.lstrip("#")
        if digits.isdigit():
            parts.append(StoreProxy.id == int(digits))
        conds.append(or_(*parts))
    return conds


def list_store_proxies(db: Session, conds: list, limit: int = 200) -> dict:
    total = db.scalar(select(func.count(StoreProxy.id)).where(*conds)) or 0
    rows = db.scalars(select(StoreProxy).where(*conds).order_by(StoreProxy.id.desc()).limit(limit)).all()
    return {"items": [store_proxy_item(p) for p in rows], "total": total, "limit": limit}


def store_summary(db: Session) -> dict:
    paid = [StoreOrder.status == "paid"]
    revenue = db.scalar(select(func.coalesce(func.sum(StoreOrder.amount_rub), 0)).where(*paid)) or 0
    paid_count = db.scalar(select(func.count(StoreOrder.id)).where(*paid)) or 0
    income = db.scalar(select(func.coalesce(func.sum(StoreOrder.income_rub), 0.0)).where(*paid)) or 0.0
    income_known = db.scalar(select(func.count(StoreOrder.id))
                             .where(*paid, StoreOrder.income_rub.isnot(None))) or 0
    cost = db.scalar(select(func.coalesce(func.sum(StoreOrder.cost_usd), 0.0))
                     .where(StoreOrder.fulfill == "done")) or 0.0
    refunded_rub = db.scalar(select(func.coalesce(func.sum(StoreOrder.amount_rub), 0))
                             .where(*paid, StoreOrder.refunded.is_(True))) or 0
    refunded_count = db.scalar(select(func.count(StoreOrder.id))
                               .where(*paid, StoreOrder.refunded.is_(True))) or 0
    by_fulfill = {f: 0 for f in STORE_FULFILL}
    for state, count in db.execute(select(StoreOrder.fulfill, func.count(StoreOrder.id))
                                   .where(StoreOrder.fulfill != "").group_by(StoreOrder.fulfill)).all():
        by_fulfill[state] = int(count)
    attention = db.scalar(select(func.count(StoreOrder.id))
                          .where(StoreOrder.fulfill.in_(ATTENTION_FULFILL), StoreOrder.refunded.is_(False))) or 0
    active = db.scalar(select(func.count(StoreProxy.id)).where(StoreProxy.status == "active")) or 0
    return {
        "revenue_rub": int(revenue),
        "paid_count": int(paid_count),
        "income_rub": round(float(income), 2),
        "income_known": int(income_known),
        "cost_usd": round(float(cost), 2),
        "refunded_rub": int(refunded_rub),
        "refunded_count": int(refunded_count),
        "by_fulfill": by_fulfill,
        "attention": int(attention),
        "active_proxies": int(active),
    }


# ---------- exports ----------

EXPORT_HEADERS = ["№", "Создан (МСК)", "Статус", "Тариф", "Сумма ₽", "К зачислению ₽", "Платёжка",
                  "Чем платили", "Причина неуспеха", "Email", "Ключ", "Оплачен (МСК)", "Устройств",
                  "Последний вход (МСК)", "ID платежа", "Тест"]


def export_rows(db: Session, conds: list) -> list[list]:
    orders = db.scalars(select(Order).where(*conds).order_by(Order.id.desc())).all()
    info = license_info_for_orders(db, orders)
    rows = []
    for o in orders:
        li = info.get(o.id) or {}
        last_seen = li.get("last_seen")
        last_seen_dt = datetime.strptime(last_seen, "%Y-%m-%dT%H:%M:%SZ") if last_seen else None
        rows.append([
            o.id,
            fmt_msk(o.created_at),
            STATUS_TEXT.get(o.status, o.status),
            plan_title(o.plan_id),
            o.amount_rub,
            o.income_rub if o.income_rub is not None else "",
            METHOD_TEXT.get(o.method, o.method),
            pay_type_text(o.pay_type) or "",
            fail_reason_text(o) or "",
            o.email,
            o.key or "",
            fmt_msk(o.paid_at),
            li.get("device_count", 0),
            fmt_msk(last_seen_dt),
            o.provider_payment_id or "",
            "да" if is_test_email(o.email) else "",
        ])
    return rows


def export_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow(EXPORT_HEADERS)
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D5BFF")
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [7, 17, 15, 11, 10, 14, 11, 17, 40, 28, 24, 17, 10, 19, 38, 6]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
